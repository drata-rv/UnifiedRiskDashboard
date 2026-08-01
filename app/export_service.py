"""
Board-ready exports: heatmap as PNG/PDF, risk register as CSV/XLSX.

Heatmap rendering reuses the exact grid math and tier colors the live page
uses (see heatmap.py, models.py) so an export always matches what's on
screen. matplotlib (Agg backend) draws both PNG and PDF from the same
figure — no headless browser, no system libraries (unlike e.g. weasyprint).
"""

import csv
import io
import json
import re

import matplotlib
matplotlib.use("Agg")
from matplotlib import pyplot as plt
from matplotlib.patches import Rectangle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, heatmap as heatmap_mod
from .models import STATUS_LABELS, TIER_COLORS, TIER_ROW_COLORS, TREATMENT_LABELS, level_label, score_to_tier

_SLUG_RE = re.compile(r"[^A-Za-z0-9]+")
_SHEET_INVALID_RE = re.compile(r"[\[\]:\*\?/\\]")


def _slug(s: str) -> str:
    return _SLUG_RE.sub("_", s or "").strip("_") or "register"


def stamped_filename(tenant_name: str, register_name: str, kind: str, ext: str) -> str:
    ts = db.now()[:16].replace("-", "").replace(":", "").replace("T", "_")
    return f"{_slug(tenant_name)}_{_slug(register_name)}_{kind}_{ts}.{ext}"


# ---------------------------------------------------------------------------
# Heatmap image export
# ---------------------------------------------------------------------------

def _draw_panel(ax, grid: list, size: int, title: str):
    for row in grid:
        for cell in row:
            x, y = cell["likelihood"] - 1, cell["impact"] - 1
            tier = score_to_tier(cell["impact"] * cell["likelihood"])
            color = f"#{TIER_COLORS.get(tier, TIER_COLORS['Unscored'])}"
            ax.add_patch(Rectangle((x, y), 1, 1, facecolor=color, edgecolor="white", linewidth=1.5))
            if cell["ids"]:
                label = "\n".join(f"#{i}" for i in cell["ids"])
                ax.text(x + 0.5, y + 0.5, label, ha="center", va="center", fontsize=8)
    ax.set_xlim(0, size)
    ax.set_ylim(0, size)
    ax.set_xticks([i + 0.5 for i in range(size)])
    ax.set_xticklabels([level_label(i + 1) for i in range(size)], fontsize=8)
    ax.set_yticks([i + 0.5 for i in range(size)])
    ax.set_yticklabels([level_label(i + 1) for i in range(size)], fontsize=8)
    ax.set_xlabel("Likelihood", fontsize=9)
    ax.set_ylabel("Impact", fontsize=9)
    ax.set_title(title, fontsize=11)
    ax.set_aspect("equal")


def build_heatmap_figure(risks: list, register_name: str, view: str = None):
    """view is None (both inherent+residual, matching the live page) or 'inherent'/'residual'."""
    size = heatmap_mod.grid_size(risks)
    seq = heatmap_mod.seq_map(risks)
    inherent_grid, _ = heatmap_mod.build_heatmap(risks, size, "impact", "likelihood", seq)
    residual_grid, _ = heatmap_mod.build_heatmap(risks, size, "residual_impact", "residual_likelihood", seq)

    panel_w = 1.5 * size + 1.8
    panel_h = 1.5 * size + 1.4
    if view is None:
        fig, axes = plt.subplots(1, 2, figsize=(panel_w * 2, panel_h))
        _draw_panel(axes[0], inherent_grid, size, "Inherent Risk")
        _draw_panel(axes[1], residual_grid, size, "Residual Risk")
    elif view == "inherent":
        fig, ax = plt.subplots(figsize=(panel_w, panel_h))
        _draw_panel(ax, inherent_grid, size, "Inherent Risk")
    elif view == "residual":
        fig, ax = plt.subplots(figsize=(panel_w, panel_h))
        _draw_panel(ax, residual_grid, size, "Residual Risk")
    else:
        raise ValueError(f"unknown view {view!r}")

    fig.suptitle(register_name, fontsize=13, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    return fig


def render_heatmap(risks: list, register_name: str, fmt: str, view: str = None) -> bytes:
    fig = build_heatmap_figure(risks, register_name, view)
    buf = io.BytesIO()
    fig.savefig(buf, format=fmt, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Risk register table export (CSV / XLSX)
# ---------------------------------------------------------------------------

_EXPORT_COLUMNS = [
    ("risk_code", "Risk Code"),
    ("title", "Risk Statement"),
    ("description", "Description"),
    ("owners", "Owners"),
    ("categories", "Categories"),
    ("impact_label", "Impact"),
    ("likelihood_label", "Likelihood"),
    ("score", "Score"),
    ("severity_tier", "Severity"),
    ("treatment_label", "Treatment"),
    ("treatment_details", "Treatment Details"),
    ("status_label", "Status"),
    ("residual_impact_label", "Residual Impact"),
    ("residual_likelihood_label", "Residual Likelihood"),
    ("residual_score", "Residual Score"),
    ("residual_tier", "Residual Severity"),
]

_COLUMN_WIDTHS = [12, 34, 40, 18, 18, 12, 12, 8, 12, 16, 34, 12, 14, 16, 14, 16]


def build_register_rows(tenant_name: str, register_id: int) -> list:
    risks = db.list_risks(tenant_name=tenant_name, register_id=register_id)
    rows = []
    for r in risks:
        owners = json.loads(r["owners_json"]) if r["owners_json"] else []
        categories = json.loads(r["categories_json"]) if r["categories_json"] else []
        rows.append({
            "risk_code": r["risk_code"] or "",
            "title": r["title"] or "",
            "description": r["description"] or "",
            "owners": ", ".join(o["name"] for o in owners),
            "categories": ", ".join(categories),
            "impact_label": level_label(r["impact"]),
            "likelihood_label": level_label(r["likelihood"]),
            "score": r["score"] if r["score"] is not None else "",
            "severity_tier": score_to_tier(r["score"]),
            "treatment_label": TREATMENT_LABELS.get(r["treatment_plan"], r["treatment_plan"] or ""),
            "treatment_details": r["treatment_details"] or "",
            "status_label": STATUS_LABELS.get(r["status"], r["status"] or ""),
            "residual_impact_label": level_label(r["residual_impact"]),
            "residual_likelihood_label": level_label(r["residual_likelihood"]),
            "residual_score": r["residual_score"] if r["residual_score"] is not None else "",
            "residual_tier": score_to_tier(r["residual_score"]),
        })
    return rows


def rows_to_csv(rows: list) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in _EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([row[key] for key, _ in _EXPORT_COLUMNS])
    return buf.getvalue().encode("utf-8")


def rows_to_xlsx(rows: list, register_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (_SHEET_INVALID_RE.sub(" ", register_name or "Risks").strip() or "Risks")[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="FF1F2937")
    for col, (_, label) in enumerate(_EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    severity_col = next(i for i, (key, _) in enumerate(_EXPORT_COLUMNS, start=1) if key == "severity_tier")
    residual_severity_col = next(i for i, (key, _) in enumerate(_EXPORT_COLUMNS, start=1) if key == "residual_tier")

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, (key, _) in enumerate(_EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row[key])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r_idx, column=severity_col).fill = PatternFill(
            "solid", fgColor="FF" + TIER_ROW_COLORS.get(row["severity_tier"], TIER_ROW_COLORS["Unscored"]))
        ws.cell(row=r_idx, column=residual_severity_col).fill = PatternFill(
            "solid", fgColor="FF" + TIER_ROW_COLORS.get(row["residual_tier"], TIER_ROW_COLORS["Unscored"]))

    for i, w in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
